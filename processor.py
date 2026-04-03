import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font as XLFont
from openpyxl.utils import get_column_letter
from fpdf import FPDF
import os

def get_processed_data(input_file):
    try:
        with open(input_file, 'rb') as f:
            header = f.read(2)
        
        if header == b'PK':
            df = pd.read_excel(input_file)
        else:
            try:
                df = pd.read_csv(input_file)
            except UnicodeDecodeError:
                df = pd.read_csv(input_file, encoding='latin1')
    except Exception as e:
        print(f"Error reading file: {e}")
        return None, None

    df.columns = [col.strip() for col in df.columns]
    
    if 'DateTime' not in df.columns:
        for col in df.columns:
            if 'date' in col.lower() and 'time' in col.lower():
                df.rename(columns={col: 'DateTime'}, inplace=True)
                break
    
    df['DateTime'] = pd.to_datetime(df['DateTime'], errors='coerce')
    df = df.dropna(subset=['DateTime'])
    
    mapping = {
        'EnNo': ['EnNo', 'Employee No', 'ID', 'Code', 'No'],
        'Name': ['Name', 'Employee Name', 'Staff Name'],
        'IOMd': ['IOMd', 'Mode', 'Direction']
    }
    
    for target, variations in mapping.items():
        if target not in df.columns:
            for v in variations:
                if v.lower() in [c.lower() for c in df.columns]:
                    actual_col = [c for c in df.columns if c.lower() == v.lower()][0]
                    df.rename(columns={actual_col: target}, inplace=True)
                    break
    
    if 'IOMd' not in df.columns:
        df['IOMd'] = 1
    
    if 'EnNo' not in df.columns or 'Name' not in df.columns:
        print(f"Error: Required columns missing. Found: {df.columns.tolist()}")
        return None, None
        
    df = df.dropna(subset=['EnNo', 'Name', 'DateTime'])
    df['Date'] = df['DateTime'].dt.date
    df['Time'] = df['DateTime'].dt.strftime('%H:%M')
    
    agg = df.groupby(['EnNo', 'Name', 'IOMd', 'Date'])['Time'].agg(['min', 'max']).reset_index()
    agg.rename(columns={'min': 'In', 'max': 'Out'}, inplace=True)
    
    # Combine In and Out for compact representation
    agg['Stacked'] = agg.apply(lambda r: f"{r['In']}\n{r['Out']}" if r['In'] != r['Out'] else f"{r['In']}", axis=1)
    
    min_date = agg['Date'].min()
    max_date = agg['Date'].max()
    start_date = min_date.replace(day=1)
    all_dates = pd.date_range(start=start_date, end=max_date).date.tolist()
    
    return agg, all_dates

def process_attendance_horizontal(agg, all_dates, output_file):
    pivot_df = agg.pivot(index=['EnNo', 'Name', 'IOMd'], columns='Date', values='Stacked')
    pivot_df = pivot_df.reindex(columns=all_dates).fillna("A")

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    
    date_range_str = f"For Period : {all_dates[0].strftime('%d-%b-%Y')} To {all_dates[-1].strftime('%d-%b-%Y')}"
    
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    date_fill = PatternFill(start_color="F79646", end_color="F79646", fill_type="solid")
    side_header_fill = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")
    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    wrap_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    bold_font = XLFont(bold=True)
    title_font = XLFont(bold=True, size=14)
    
    # Report Headers (Rows 1 and 2)
    total_cols = 3 + len(all_dates)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    cell1 = ws.cell(row=1, column=1, value="Monthly Attendance Report with (In\\Out) Time")
    cell1.font = title_font; cell1.alignment = Alignment(horizontal='center')
    
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    cell2 = ws.cell(row=2, column=1, value=date_range_str)
    cell2.font = bold_font; cell2.alignment = Alignment(horizontal='center')
    
    # Table Headers (Row 3)
    ws.cell(row=3, column=1, value="Eno").fill = side_header_fill
    ws.cell(row=3, column=2, value="Name").fill = side_header_fill
    ws.cell(row=3, column=3, value="IOMd").fill = side_header_fill
    
    col_idx = 4
    for date in all_dates:
        cell = ws.cell(row=3, column=col_idx, value=date.strftime('%d-%b'))
        cell.fill = date_fill; cell.font = bold_font; cell.alignment = wrap_center
        col_idx += 1
        
    for c in range(1, 4):
        ws.cell(row=3, column=c).font = bold_font; ws.cell(row=3, column=c).alignment = wrap_center

    # Data Rows (Starting Row 4)
    row_idx = 4
    for user_info, data in pivot_df.iterrows():
        ws.cell(row=row_idx, column=1, value=str(user_info[0]))
        ws.cell(row=row_idx, column=2, value=user_info[1])
        ws.cell(row=row_idx, column=3, value=user_info[2])
        col_idx = 4
        for date in all_dates:
            val = data.get(date, "A")
            ws.cell(row=row_idx, column=col_idx, value=val)
            ws.cell(row=row_idx, column=col_idx).alignment = wrap_center
            col_idx += 1
        row_idx += 1
        
    for row in ws.iter_rows(min_row=3, max_row=row_idx-1, min_col=1, max_col=total_cols):
        for cell in row: cell.border = thin_border
    
    ws.column_dimensions['A'].width = 8; ws.column_dimensions['B'].width = 25; ws.column_dimensions['C'].width = 8
    for i in range(4, total_cols + 1): ws.column_dimensions[get_column_letter(i)].width = 10
        
    wb.save(output_file)

def process_attendance_vertical(agg, all_dates, output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vertical Attendance"
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    sub_header_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
    side_header_fill = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    bold_font = XLFont(bold=True); white_bold_font = XLFont(bold=True, color="FFFFFF"); center_align = Alignment(horizontal='center', vertical='center')
    
    # Global Header
    ws.merge_cells('A1:C1')
    ws.cell(row=1, column=1, value="Monthly Attendance Report with (In\\Out) Time").font = XLFont(bold=True, size=12)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal='center')
    
    row_idx = 3
    teachers = agg.groupby(['EnNo', 'Name', 'IOMd'])
    for (eno, name, iomd), t_data in teachers:
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
        cell = ws.cell(row=row_idx, column=1, value=f"Teacher: {name} (ID: {eno})")
        cell.fill = header_fill; cell.font = white_bold_font; cell.alignment = center_align
        row_idx += 1
        
        ws.cell(row=row_idx, column=1, value="Date").fill = side_header_fill
        ws.cell(row=row_idx, column=2, value="In").fill = sub_header_fill
        ws.cell(row=row_idx, column=3, value="Out").fill = sub_header_fill
        for c in range(1, 4): ws.cell(row=row_idx, column=c).font = bold_font
        row_idx += 1
        
        t_data_indexed = t_data.set_index('Date')
        for date in all_dates:
            ws.cell(row=row_idx, column=1, value=date.strftime('%d-%b-%y'))
            if date in t_data_indexed.index:
                ws.cell(row=row_idx, column=2, value=t_data_indexed.loc[date, 'In'])
                ws.cell(row=row_idx, column=3, value=t_data_indexed.loc[date, 'Out'])
            else:
                ws.cell(row=row_idx, column=2, value="A"); ws.cell(row=row_idx, column=3, value="A")
            for c in range(1, 4): ws.cell(row=row_idx, column=c).border = thin_border
            row_idx += 1
        row_idx += 2
        
    for i in range(1, 4): ws.column_dimensions[get_column_letter(i)].width = 25
    wb.save(output_file)

def process_to_pdf_horizontal(agg, all_dates, output_file, layout="paginated"):
    CHUNK_SIZE = 12 if layout == "paginated" else len(all_dates)
    date_chunks = [all_dates[i:i + CHUNK_SIZE] for i in range(0, len(all_dates), CHUNK_SIZE)]
    pdf = FPDF(orientation='L', unit='mm', format='A4')
    c_header, c_date, c_sub_header, c_side_header = (79, 129, 189), (247, 150, 70), (198, 224, 180), (253, 233, 217)
    
    pivot_df = agg.pivot(index=['EnNo', 'Name', 'IOMd'], columns='Date', values='Stacked')
    pivot_df = pivot_df.reindex(columns=all_dates).fillna("A")

    for chunk in date_chunks:
        pdf.add_page()
        pdf.set_font("Helvetica", 'B', size=14)
        pdf.cell(0, 8, "Monthly Attendance Report with (In\\Out) Time", ln=True, align='C')
        pdf.set_font("Helvetica", 'B', size=11)
        pdf.cell(0, 8, f"For Period : {chunk[0].strftime('%d-%b-%Y')} To {chunk[-1].strftime('%d-%b-%Y')}", ln=True, align='C')
        pdf.ln(2)
        
        num_dates = len(chunk); avail_width = 277
        if layout == "single":
            w_eno, w_name, w_iomd = 10, 35, 8; f_size = 6.5
        else:
            w_eno, w_name, w_iomd = 12, 45, 12; f_size = 8
            
        w_date_col = (avail_width - (w_eno + w_name + w_iomd)) / num_dates
        
        pdf.set_fill_color(*c_side_header); pdf.set_text_color(0); pdf.set_font("Helvetica", 'B', size=f_size)
        pdf.cell(w_eno, 10, "Eno", border=1, align='C', fill=True)
        pdf.cell(w_name, 10, "Emp Name", border=1, align='C', fill=True)
        pdf.cell(w_iomd, 10, "IOMd", border=1, align='C', fill=True)
        pdf.set_fill_color(*c_date)
        for date in chunk:
            pdf.cell(w_date_col, 10, date.strftime('%d-%m'), border=1, align='C', fill=True)
        pdf.ln()

        pdf.set_font("Helvetica", size=f_size - 0.5)
        for user_info, data in pivot_df.iterrows():
            row_h = 8 if layout == "paginated" else 6
            curr_x, curr_y = pdf.get_x(), pdf.get_y()
            pdf.cell(w_eno, row_h, str(user_info[0]), border=1, align='C')
            pdf.cell(w_name, row_h, str(user_info[1])[:18], border=1)
            pdf.cell(w_iomd, row_h, str(user_info[2]), border=1, align='C')
            for date in chunk:
                val = data.get(date, "A")
                tx, ty = pdf.get_x(), pdf.get_y()
                pdf.multi_cell(w_date_col, row_h/2 if "\n" in str(val) else row_h, str(val), border=1, align='C')
                pdf.set_xy(tx + w_date_col, ty)
            pdf.ln(row_h)
    pdf.output(output_file)

def process_to_pdf_vertical(agg, all_dates, output_file):
    pdf = FPDF(orientation='P', unit='mm', format='A4')
    c_header, c_sub_header, c_side_header = (79, 129, 189), (198, 224, 180), (253, 233, 217)
    teachers = agg.groupby(['EnNo', 'Name', 'IOMd'])
    for (eno, name, iomd), t_data in teachers:
        pdf.add_page()
        pdf.set_font("Helvetica", 'B', size=14)
        pdf.cell(0, 10, "Monthly Attendance Report with (In\\Out) Time", ln=True, align='C')
        pdf.ln(5)
        pdf.set_fill_color(*c_header); pdf.set_text_color(255); pdf.set_font("Helvetica", 'B', size=12)
        pdf.cell(0, 10, f"Teacher: {name} (ID: {eno})", border=1, align='C', fill=True)
        pdf.ln(12)
        pdf.set_fill_color(*c_side_header); pdf.set_text_color(0); pdf.set_font("Helvetica", 'B', size=10)
        pdf.cell(60, 10, "Date", border=1, align='C', fill=True)
        pdf.set_fill_color(*c_sub_header)
        pdf.cell(65, 10, "In Time", border=1, align='C', fill=True)
        pdf.cell(65, 10, "Out Time", border=1, align='C', fill=True); pdf.ln()
        pdf.set_font("Helvetica", size=10); t_data_indexed = t_data.set_index('Date')
        for date in all_dates:
            pdf.cell(60, 8, date.strftime('%d-%b-%Y'), border=1, align='C')
            if date in t_data_indexed.index:
                pdf.cell(65, 8, str(t_data_indexed.loc[date, 'In']), border=1, align='C')
                pdf.cell(65, 8, str(t_data_indexed.loc[date, 'Out']), border=1, align='C')
            else:
                pdf.cell(65, 8, "A", border=1, align='C'); pdf.cell(65, 8, "A", border=1, align='C')
            pdf.ln()
    pdf.output(output_file)

def main_process(input_file, export_type, layout="horizontal"):
    agg, all_dates = get_processed_data(input_file)
    if agg is None: return
    output_file = "attendance_output.xlsx" if export_type == "excel" else "attendance_output.pdf"
    try:
        if export_type == "excel":
            if layout == "horizontal": process_attendance_horizontal(agg, all_dates, output_file)
            else: process_attendance_vertical(agg, all_dates, output_file)
        else:
            if layout == "horizontal": process_to_pdf_horizontal(agg, all_dates, output_file, layout="paginated")
            elif layout == "single": process_to_pdf_horizontal(agg, all_dates, output_file, layout="single")
            else: process_to_pdf_vertical(agg, all_dates, output_file)
        print(f"File saved to {output_file}")
    except Exception as e: print(f"Error processing {export_type}: {e}"); raise e

if __name__ == "__main__":
    main_process('log_input.csv', 'excel', 'horizontal')
    main_process('log_input.csv', 'pdf', 'horizontal')
